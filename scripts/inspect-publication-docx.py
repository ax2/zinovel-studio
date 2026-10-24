"""Independent OOXML/editability inspection and optional native-reader PDF raster QA."""
import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
import pypdf
import pypdfium2 as pdfium
import pdfplumber

parser = argparse.ArgumentParser()
parser.add_argument("docx")
parser.add_argument("expected")
parser.add_argument("output")
parser.add_argument("--pdf")
args = parser.parse_args()
source, out = Path(args.docx), Path(args.output)
out.mkdir(parents=True, exist_ok=True)
expected = json.loads(Path(args.expected).read_text(encoding="utf-8-sig"))
namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
with ZipFile(source) as archive:
    assert archive.testzip() is None
    names = archive.namelist()
    assert len(names) == len(set(names))
    assert not any(name.startswith("/") or "\\" in name or ".." in name.split("/") for name in names)
    assert not any("vba" in name.lower() or "/embeddings/" in name.lower() for name in names)
    for name in names:
        if name.endswith(".xml") or name.endswith(".rels"):
            root = ET.fromstring(archive.read(name))
            instructions = [(item.text or "") for item in root.findall(".//w:instrText", namespaces)]
            instructions += [item.attrib.get(f"{{{namespaces['w']}}}instr", "") for item in root.findall(".//w:fldSimple", namespaces)]
            assert all(value.strip() in {"PAGE", "NUMPAGES"} for value in instructions)
            if name.endswith(".rels"):
                assert all(item.attrib.get("TargetMode") != "External" for item in root)
    xml = ET.fromstring(archive.read("word/document.xml"))
    style_xml = ET.fromstring(archive.read("word/styles.xml"))
    defaults = style_xml.find("w:docDefaults", namespaces)
    if "size" in expected:
        assert defaults.find("w:rPrDefault/w:rPr/w:sz", namespaces).attrib[f"{{{namespaces['w']}}}val"] == str(expected["size"] * 2)
    if "font" in expected:
        assert defaults.find("w:rPrDefault/w:rPr/w:rFonts", namespaces).attrib[f"{{{namespaces['w']}}}eastAsia"] == {"serif": "SimSun", "sans": "Microsoft YaHei"}[expected["font"]]
    if "leading" in expected:
        assert defaults.find("w:pPrDefault/w:pPr/w:spacing", namespaces).attrib[f"{{{namespaces['w']}}}line"] == str(round(expected["leading"] * 240))
    styles = [item.attrib.get(f"{{{namespaces['w']}}}val") for item in xml.findall(".//w:pStyle", namespaces)]
    assert styles[0] == "Title"
    assert styles.count("Heading1") == expected["chapters"]
    for field in xml.findall(".//w:instrText", namespaces):
        assert (field.text or "").strip() in {"PAGE", "NUMPAGES"}
document = Document(source)
text = "\n".join(paragraph.text for paragraph in document.paragraphs)
compact = "".join(text.split())
for value in expected["required"]:
    assert "".join(value.split()) in compact, f"Missing editable content: {value!r}"
for value in expected["forbidden"]:
    assert "".join(value.split()) not in compact, f"Excluded content: {value!r}"
position = -1
for value in expected["ordered"]:
    found = compact.find("".join(value.split()), position + 1)
    assert found > position
    position = found
for value, count in expected.get("counts", {}).items():
    assert text.count(value) == count, f"Character count: {value!r}"
dimensions = {"Letter": (8.5, 11), "A4": (210 / 25.4, 297 / 25.4), "A5": (148 / 25.4, 210 / 25.4)}[expected["paper"]]
for section in document.sections:
    assert abs(section.page_width.inches - dimensions[0]) < 0.01
    assert abs(section.page_height.inches - dimensions[1]) < 0.01
report = {"status": "passed", "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
          "paragraphs": len(document.paragraphs), "chapterHeadings": styles.count("Heading1"),
          "paper": expected["paper"], "allTextEditable": True,
          "scope": "Independent python-docx and OOXML checks; optional PDF comes from a separately identified native reader"}
if args.pdf:
    pdf_source = Path(args.pdf)
    reader = pypdf.PdfReader(pdf_source)
    for page in reader.pages:
        assert abs(float(page.mediabox.width) - dimensions[0] * 72) < 1
        assert abs(float(page.mediabox.height) - dimensions[1] * 72) < 1
    with pdfplumber.open(pdf_source) as spatial:
        body_texts = []
        footer_texts = []
        # Separate only the known footer band, and verify its complete content.
        # This preserves a title or paragraph continued across a page boundary.
        bottom_margin = document.sections[0].bottom_margin.pt
        for index, page in enumerate(spatial.pages):
            boundary = page.height - bottom_margin
            footer = page.crop((0, boundary, page.width, page.height)).extract_text() or ""
            assert "".join(footer.split()) == f"{index + 1}/{len(spatial.pages)}", f"Unexpected footer: {footer!r}"
            footer_texts.append(footer)
            body_texts.append(page.crop((0, 0, page.width, boundary)).extract_text() or "")
        raw_pdf_text = "\n".join(body_texts)
    pdf_text = "".join(raw_pdf_text.split())
    for value in expected["required"]:
        assert "".join(value.split()) in pdf_text, f"Missing native rendered content: {value!r}"
    for value in expected["forbidden"]:
        assert "".join(value.split()) not in pdf_text
    assert "\ufffd" not in pdf_text
    for value, count in expected.get("counts", {}).items():
        assert raw_pdf_text.count(value) == count, f"Native rendered character count: {value!r}"
    with pdfplumber.open(pdf_source) as parsed:
        outside = [{"page": index + 1, "text": char["text"]} for index, page in enumerate(parsed.pages)
                   for char in page.chars if char["text"].strip() and
                   (char["x0"] < -1 or char["x1"] > page.width + 1 or char["top"] < -1 or char["bottom"] > page.height + 1)]
        assert not outside, f"Native glyphs outside paper: {outside[:5]}"
    pdf = pdfium.PdfDocument(str(pdf_source))
    images = []
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            bitmap = page.render(scale=1.5)
            name = f"page-{index + 1:03}.png"
            bitmap.to_pil().save(out / name)
            images.append(name)
            bitmap.close()
            page.close()
    finally:
        pdf.close()
    report.update({"pdfPages": len(reader.pages), "pdfSha256": hashlib.sha256(pdf_source.read_bytes()).hexdigest(), "renderedPages": images, "verifiedFooters": footer_texts})
(out / "docx-inspection.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(report, ensure_ascii=False))
