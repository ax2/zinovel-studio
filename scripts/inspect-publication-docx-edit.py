"""Inspect the saved synthetic Word copy without using the production generator."""
import argparse
import hashlib
import json
from pathlib import Path
from docx import Document

parser = argparse.ArgumentParser()
parser.add_argument("original")
parser.add_argument("edited")
parser.add_argument("expected_sha256")
parser.add_argument("output")
parser.add_argument("--pdf")
args = parser.parse_args()
original, edited = Path(args.original), Path(args.edited)
digest = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
assert digest(original) == args.expected_sha256
before, after = Document(original), Document(edited)
original_paragraphs = [paragraph.text for paragraph in before.paragraphs]
edited_paragraphs = [paragraph.text for paragraph in after.paragraphs]
assert edited_paragraphs[:len(original_paragraphs)] == original_paragraphs
assert edited_paragraphs[len(original_paragraphs):] == ["Word editable verification marker 𠮷 🌊。"]
report = {
    "status": "passed", "originalSha256": digest(original), "editedSha256": digest(edited),
    "paragraphsBefore": len(original_paragraphs), "paragraphsAfter": len(edited_paragraphs),
    "originalUnchanged": True, "originalParagraphPrefixUnchanged": True,
    "appendedMarkerExact": True, "scope": "Independent python-docx inspection of an actual native Word edit/save; synthetic content only",
}
if args.pdf:
    import pdfplumber
    body = []
    footers = []
    with pdfplumber.open(args.pdf) as pdf:
        for index, page in enumerate(pdf.pages):
            boundary = page.height - 0.55 * 72
            body.append(page.crop((0, 0, page.width, boundary)).extract_text() or "")
            footer = "".join((page.crop((0, boundary, page.width, page.height)).extract_text() or "").split())
            assert footer == f"{index + 1}/{len(pdf.pages)}"
            footers.append(footer)
    expected_text = "".join("".join(text.split()) for text in edited_paragraphs)
    actual_text = "".join("".join(text.split()) for text in body)
    assert actual_text == expected_text, "Word PDF text does not match the complete saved editable text"
    report.update({"pdfSha256": digest(Path(args.pdf)), "renderedTextMatchesAllEditedParagraphs": True, "verifiedFooters": footers})
Path(args.output).write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(report, ensure_ascii=False))
